#!/usr/bin/env python3
"""Build the Dockerfile Audit work report in PDF and DOCX with embedded proofs."""
import os
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
from reportlab.lib.enums import TA_LEFT
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = os.path.expanduser("~/workspace/dockerfile-audit")
SCREEN = os.path.join(OUT, "screenshots", "size-comparison.png")
REPO = "https://github.com/pakmultilinks-dot/dockerfile-audit"

DARK = HexColor("#1a1a2e")
ACCENT = HexColor("#0f62fe")
GRAY = HexColor("#555555")

styles = getSampleStyleSheet()
title_s = ParagraphStyle("Title2", parent=styles["Title"], fontSize=22, textColor=DARK, spaceAfter=4*mm)
h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=14, textColor=DARK, spaceBefore=6*mm, spaceAfter=3*mm)
h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=11, textColor=DARK, spaceBefore=4*mm, spaceAfter=2*mm)
body = ParagraphStyle("Body2", parent=styles["BodyText"], fontSize=10, leading=14, spaceAfter=2*mm)
small = ParagraphStyle("Small", parent=styles["BodyText"], fontSize=9, leading=12, textColor=GRAY)
code_s = ParagraphStyle("Code", parent=styles["Code"], fontSize=8.5, leading=11, fontName="Courier")

findings = [
    ("P1 / N1", "Base image uses :latest tag", "Unpinned base; builds are not reproducible and can break or change silently.", "High"),
    ("P2", "apt-get without cleanup", "Leaves ~100MB+ of package lists in the layer.", "Medium"),
    ("P3", "Unnecessary build tools in runtime (gcc, curl)", "Bloats image and widens attack surface.", "Medium"),
    ("P4", "pip install without --no-cache-dir", "Pip cache bloats the layer.", "Low"),
    ("P5", "COPY . . before dependency install", "Any code change invalidates the dependency layer cache; slower rebuilds.", "Medium"),
    ("P6", "Missing .dockerignore", "Build context ships .git, __pycache__, node_modules to the daemon.", "Medium"),
    ("P7", "Container runs as root", "Security risk; violates least privilege.", "High"),
    ("P8", "Shell-form CMD", "PID 1 is sh, signals do not reach the app; graceful shutdown broken.", "Medium"),
    ("P9", "No HEALTHCHECK", "Orchestrator cannot detect a hung app.", "Low"),
    ("P10", "Hardcoded secrets in ENV", "API keys baked into image layers; leaked to anyone with the image.", "Critical"),
    ("N2", "npm install instead of npm ci", "Non-reproducible dependency tree.", "Medium"),
    ("N3", "Dev dependencies shipped to production", "Larger image, larger attack surface.", "Medium"),
    ("N4", "No NODE_ENV=production", "Slower runtime, verbose warnings.", "Low"),
    ("G1", "Single-stage Go build ships the toolchain", "Go toolchain (~800MB) in the runtime image.", "High"),
    ("G2", "Binary built without -ldflags='-s -w'", "Debug symbols bloat the binary.", "Low"),
    ("G3", "go.mod/go.sum copied after source", "Dependency layer never cached.", "Medium"),
]

sizes = [
    ["Sample", "Before (measured)", "After (measured)", "Reduction"],
    ["app-python (Flask)", "1.2 GB", "136 MB", "88.7%"],
    ["app-node (Express)", "1.25 GB", "230 MB", "81.6%"],
    ["app-go (net/http)", "893 MB", "81.9 MB", "90.8%"],
]

def build_pdf():
    path = os.path.join(OUT, "Dockerfile-Audit-Report.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=16*mm, bottomMargin=16*mm)
    st = []
    st.append(Paragraph("Dockerfile Audit Report", title_s))
    st.append(Paragraph("TechAbout DevOps task ZR-26-00755 (Task 4 of 20) - Abdur Rehman - 19 Oct 2026", small))
    st.append(Spacer(1, 4*mm))

    st.append(Paragraph("1. What was done", h1))
    st.append(Paragraph(
        "Audited 3 sample Dockerfiles (Python/Flask, Node.js/Express, Go net/http). "
        "Found 16 problems across the three files, wrote fixed versions of each, "
        "and measured the real image sizes before and after. All work is in the GitHub repo below.", body))

    st.append(Paragraph("2. Findings", h1))
    st.append(Paragraph("16 issues found: 1 Critical, 3 High, 8 Medium, 4 Low. Full detail with fixes in docs/AUDIT.md.", body))
    tdata = [["ID", "Issue", "Why it matters", "Severity"]] + [list(f) for f in findings]
    t = Table(tdata, colWidths=[18*mm, 52*mm, 62*mm, 22*mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), DARK),
        ("TEXTCOLOR", (0, 0), (-1, 0), HexColor("#ffffff")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEADING", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.4, HexColor("#cccccc")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#ffffff"), HexColor("#f4f6fb")]),
    ]))
    # wrap cells as paragraphs for wrapping
    wrapped = []
    for r, row in enumerate(tdata):
        wrapped.append([Paragraph(f"<b>{c}</b>" if r == 0 else str(c), code_s if r > 0 and i == 0 else body) for i, c in enumerate(row)])
    t = Table(wrapped, colWidths=[16*mm, 50*mm, 64*mm, 22*mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), DARK),
        ("TEXTCOLOR", (0, 0), (-1, 0), HexColor("#ffffff")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEADING", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.4, HexColor("#cccccc")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#ffffff"), HexColor("#f4f6fb")]),
    ]))
    st.append(t)

    st.append(Paragraph("3. Size comparison (measured with docker images)", h1))
    st2 = [[Paragraph(f"<b>{c}</b>", body) for c in sizes[0]]] + [[Paragraph(c, body) for c in row] for row in sizes[1:]]
    t2 = Table(st2, colWidths=[45*mm, 38*mm, 38*mm, 33*mm])
    t2.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
        ("TEXTCOLOR", (0, 0), (-1, 0), HexColor("#ffffff")),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.5, HexColor("#999999")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#ffffff"), HexColor("#eef4ff")]),
    ]))
    st.append(t2)
    st.append(Spacer(1, 3*mm))
    st.append(Paragraph("Proof: terminal output of the measurement run:", h2))
    if os.path.exists(SCREEN):
        st.append(Image(SCREEN, width=150*mm, height=150*mm*570/760*0.62))
    st.append(Spacer(1, 3*mm))
    st.append(Paragraph(
        "How the numbers were produced: base images were pulled from Docker Hub and the before/after "
        "image contents were assembled layer-by-layer from the exact dependency sets each Dockerfile installs "
        "(pip install flask, npm ci, go build, apt gcc+curl, app sources), then loaded with docker load and "
        "measured with docker images. The sandbox blocks container networking, so RUN steps could not execute "
        "inside docker build; every file in the measured images is the real artifact the Dockerfile would install.", small))

    st.append(Paragraph("4. Key fixes applied", h1))
    for b in [
        "<b>Pinned slim bases:</b> python:latest -&gt; python:3.12-slim-bookworm, node:latest -&gt; node:22-bookworm-slim.",
        "<b>Multi-stage Go build:</b> toolchain compiles in the build stage; runtime is debian:12-slim with only the static binary.",
        "<b>Non-root everywhere:</b> dedicated appuser (uid 10001); images no longer run as root.",
        "<b>Cache-friendly ordering:</b> dependency manifests copied and installed before application code.",
        "<b>Exec-form CMD/ENTRYPOINT</b> plus HEALTHCHECK on all three services.",
        "<b>Secrets removed</b> from image layers (documented in AUDIT.md with the safe pattern).",
        "<b>.dockerignore</b> added to all three samples.",
    ]:
        st.append(Paragraph(b, body, bulletText="-"))

    st.append(Paragraph("5. Deliverables", h1))
    st.append(Paragraph(f"GitHub repo (public): {REPO}", body))
    for b in [
        "samples/app-python, samples/app-node, samples/app-go: Dockerfile.before / Dockerfile.after, app source, .dockerignore",
        "docs/AUDIT.md: every finding, why it matters, the exact fix",
        "docs/BEST_PRACTICES.md: 12-rule shortlist for the team",
        "scripts/build-compare.sh: rebuild all 6 images and print the size table on any machine with Docker",
        "scripts/fetch-images.py, scripts/assemble-image.py: tooling used for this audit",
    ]:
        st.append(Paragraph(b, body, bulletText="-"))

    doc.build(st)
    print("PDF:", path)

def build_docx():
    path = os.path.join(OUT, "Dockerfile-Audit-Report.docx")
    d = Document()
    style = d.styles["Normal"]
    style.font.size = Pt(10.5)
    d.add_heading("Dockerfile Audit Report", 0)
    d.add_paragraph("TechAbout DevOps task ZR-26-00755 (Task 4 of 20) - Abdur Rehman - 19 Oct 2026").italic = True

    d.add_heading("1. What was done", 1)
    d.add_paragraph(
        "Audited 3 sample Dockerfiles (Python/Flask, Node.js/Express, Go net/http). "
        "Found 16 problems across the three files, wrote fixed versions of each, "
        "and measured the real image sizes before and after. All work is in the GitHub repo: " + REPO)

    d.add_heading("2. Findings", 1)
    d.add_paragraph("16 issues found: 1 Critical, 3 High, 8 Medium, 4 Low. Full detail with fixes in docs/AUDIT.md.")
    t = d.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    for i, h in enumerate(["ID", "Issue", "Why it matters", "Severity"]):
        hdr[i].text = h
    for f in findings:
        row = t.add_row().cells
        for i, c in enumerate(f):
            row[i].text = c
            row[i].paragraphs[0].runs[0].font.size = Pt(9)

    d.add_heading("3. Size comparison (measured with docker images)", 1)
    t2 = d.add_table(rows=1, cols=4)
    t2.style = "Light Grid Accent 1"
    hdr = t2.rows[0].cells
    for i, h in enumerate(sizes[0]):
        hdr[i].text = h
    for row in sizes[1:]:
        cells = t2.add_row().cells
        for i, c in enumerate(row):
            cells[i].text = c
    d.add_paragraph("Proof: terminal output of the measurement run:")
    if os.path.exists(SCREEN):
        d.add_picture(SCREEN, width=Inches(6))
    d.add_paragraph(
        "How the numbers were produced: base images were pulled from Docker Hub and the before/after "
        "image contents were assembled layer-by-layer from the exact dependency sets each Dockerfile installs "
        "(pip install flask, npm ci, go build, apt gcc+curl, app sources), then loaded with docker load and "
        "measured with docker images. The sandbox blocks container networking, so RUN steps could not execute "
        "inside docker build; every file in the measured images is the real artifact the Dockerfile would install.").italic = True

    d.add_heading("4. Key fixes applied", 1)
    for b in [
        "Pinned slim bases: python:latest -> python:3.12-slim-bookworm, node:latest -> node:22-bookworm-slim.",
        "Multi-stage Go build: toolchain compiles in the build stage; runtime is debian:12-slim with only the static binary.",
        "Non-root everywhere: dedicated appuser (uid 10001); images no longer run as root.",
        "Cache-friendly ordering: dependency manifests copied and installed before application code.",
        "Exec-form CMD/ENTRYPOINT plus HEALTHCHECK on all three services.",
        "Secrets removed from image layers (documented in AUDIT.md with the safe pattern).",
        ".dockerignore added to all three samples.",
    ]:
        d.add_paragraph(b, style="List Bullet")

    d.add_heading("5. Deliverables", 1)
    d.add_paragraph("GitHub repo (public): " + REPO)
    for b in [
        "samples/app-python, samples/app-node, samples/app-go: Dockerfile.before / Dockerfile.after, app source, .dockerignore",
        "docs/AUDIT.md: every finding, why it matters, the exact fix",
        "docs/BEST_PRACTICES.md: 12-rule shortlist for the team",
        "scripts/build-compare.sh: rebuild all 6 images and print the size table on any machine with Docker",
        "scripts/fetch-images.py, scripts/assemble-image.py: tooling used for this audit",
    ]:
        d.add_paragraph(b, style="List Bullet")

    d.save(path)
    print("DOCX:", path)

if __name__ == "__main__":
    build_pdf()
    build_docx()
